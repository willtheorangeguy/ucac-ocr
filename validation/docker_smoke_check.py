"""Live Docker browser/import/OCR/export check with the consented reference photo.

Creates a clearly labelled test source in the running app. Requires Playwright
and an explicit local Chromium executable. Uses the configured LAN OCR server.
"""
import json
import sys
import time
from pathlib import Path

import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[1]
HOST='http://127.0.0.1:8000'
OUT=ROOT/'validation'
PHOTO=ROOT/'tests/fixtures/morning-2026-10-09.jpg'
TEMPLATE=ROOT/'Head Counts - FALL 2026.xlsx'
if not TEMPLATE.exists():TEMPLATE=ROOT/'app/assets/monthly-template.xlsx'
report=dict(host=HOST,template=TEMPLATE.name,checks={},started=time.time())


def save_report():
    (OUT/'docker-smoke-results.json').write_text(json.dumps(report,indent=2))


with httpx.Client(base_url=HOST,timeout=30,trust_env=False) as client:
    health=client.get('/health');health.raise_for_status()
    report['checks']['web_health']=health.json()
    diagnostics=client.get('/api/settings/diagnostics');diagnostics.raise_for_status()
    report['checks']['ocr_connection']=diagnostics.json()
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path=sys.argv[1],headless=True)
        page=browser.new_page(viewport=dict(width=1100,height=1000))
        errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
        try:
            if len(sys.argv)>2:
                page.goto(HOST+'/jobs/'+sys.argv[2])
            else:
                page.goto(HOST)
                assert page.locator('section.step').count()==4
                page.locator('input[name=month]').fill('2026-10')
                page.locator('input[name=template]').set_input_files(str(TEMPLATE))
                page.locator('input[name=photos]').set_input_files(dict(
                    name='SMOKE TEST - morning-2026-10-09.jpg',mimeType='image/jpeg',buffer=PHOTO.read_bytes()))
                page.get_by_role('button',name='Create batch').click()
                page.wait_for_url('**/jobs/*',timeout=60000)
            job=page.url.rsplit('/',1)[-1];report['job']=job
            report['job_url']=page.url;save_report()
            print('Created labelled test import:',page.url,flush=True)
            page.wait_for_selector('#corner-canvas',state='visible',timeout=45000)
            page.wait_for_function("document.getElementById('corner-canvas').width===1080")
            page.locator('#corner-layout').select_option('legacy')
            canvas=page.locator('#corner-canvas');box=canvas.bounding_box()
            for x,y in [[154,262],[918,252],[968,1185],[138,1200]]:
                canvas.click(position=dict(x=x*box['width']/1080,y=y*box['height']/1440))
            page.wait_for_function("document.getElementById('corner-editor').hidden",timeout=45000)
            state=client.get('/api/jobs/'+job).json()
            geometry=json.loads(state['sources'][0]['geometry'])
            assert geometry[0]['section']=='morning' and geometry[0]['layout']=='legacy'
            report['checks']['browser_corner_save']=True;save_report()
            print('Real browser upload and corner save passed; waiting for OCR.',flush=True)
            deadline=time.monotonic()+900;last=None
            while time.monotonic()<deadline:
                response=client.get('/api/jobs/'+job);response.raise_for_status();state=response.json()
                counts={name:sum(t['state']==name for t in state['tasks']) for name in ['pending','in_flight','succeeded','failed','retry_wait']}
                progress=(state['state'],counts)
                if progress!=last:
                    print('Worker:',state['state'],counts,flush=True);last=progress
                report['state']=state['state'];report['task_states']=counts;save_report()
                if state['state']=='paused':raise AssertionError(state.get('diagnostic'))
                if any(i['kind']=='geometry' for i in state['issues']):raise AssertionError('Worker rejected browser-selected geometry')
                settled=state['sources'][0]['state'] in ('succeeded','deferred','needs_review','failed')
                if settled and state['exports'] and state['exports'][0]['revision']==state['revision']:break
                time.sleep(2)
            else:raise AssertionError('Live OCR did not finish within 15 minutes')
            assert not counts['failed'] and counts['succeeded']==4,counts
            first={c['column_name']:c for c in state['cells'] if c['date']=='2026-10-09' and c['time']=='05:15'}
            actual={c:cell['value'] for c,cell in first.items()}
            report['first_row']=actual
            report['first_row_status']={c:cell['status'] for c,cell in first.items()}
            assert actual==dict(public=0,group=32,dry=2,external=None),actual
            assert all(c['status']=='accepted' for c in first.values())
            latest=state['exports'][0];export_report=json.loads(latest['report'])
            assert export_report['layout']=='monthly'
            page.reload();page.wait_for_selector('.download')
            with page.expect_download() as download:page.locator('.download').click()
            result=download.value
            assert result.suggested_filename=='2026-10-headcounts.xlsx'
            target=OUT/'docker-smoke-export.xlsx';result.save_as(target)
            workbook=load_workbook(target,read_only=True)
            try:
                sheet=workbook['DATA ENTRY - OCT 2026']
                assert [sheet[c+'564'].value for c in 'DEFG']==[0,None,32,2]
                assert sheet['G562'].value.date().isoformat()=='2026-10-09'
                assert 'MASTER' not in workbook.sheetnames
            finally:workbook.close()
            report['checks']['live_ocr_first_row']=True
            report['checks']['monthly_download']=True
            report['issue_count']=len(state['issues'])
            (OUT/'screenshots').mkdir(exist_ok=True)
            page.screenshot(path=str(OUT/'screenshots/docker-live-import.png'),full_page=True)
            page.set_viewport_size(dict(width=390,height=844))
            assert page.locator('.download').is_visible()
            page.screenshot(path=str(OUT/'screenshots/docker-live-import-mobile.png'),full_page=True)
            assert not errors,errors
            report['checks']['desktop_mobile_browser']=True
            report['checks']['no_browser_errors']=True
            report['finished']=time.time();save_report()
            print('PASS: live OCR returned Public 0 / External blank / Internal 32 / Dry 2.',flush=True)
            print('PASS: monthly workbook downloaded and verified; desktop/mobile pages have no script errors.',flush=True)
        except BaseException as e:
            report['error']=str(e);save_report();raise
        finally:browser.close()
