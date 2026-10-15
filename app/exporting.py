import json
import subprocess
import tempfile
import time
from pathlib import Path
from filelock import FileLock
from .config import settings
from .db import transaction, query
from .storage import sha, atomic_write
from .workbook import generate


def export_job(job_id):
    lockpath=settings.data_dir/job_id/'export.lock'; lockpath.parent.mkdir(parents=True,exist_ok=True)
    with FileLock(lockpath):
        with transaction() as db:
            job=dict(db.execute('SELECT * FROM jobs WHERE id=?',(job_id,)).fetchone())
            existing=db.execute('SELECT * FROM exports WHERE job_id=? AND revision=?',(job_id,job['revision'])).fetchone()
            if existing:return dict(existing)
            if sha(Path(job['template']).read_bytes())!=job['template_hash']:raise ValueError('Immutable template checksum changed')
            records=[dict(r) for r in db.execute('SELECT * FROM cells WHERE job_id=?',(job_id,))]
            source_report=[dict(r) for r in db.execute('SELECT filename,state,assigned_date,observed_date,hash FROM sources WHERE job_id=?',(job_id,))]
            issues=[dict(r) for r in db.execute('SELECT * FROM issues WHERE job_id=? AND resolved=0',(job_id,))]
            unresolved={i['cell_key'] for i in issues if i['cell_key']}
            records=[{**r,'status':'needs_review','value':None} if r['key'] in unresolved else r for r in records]
            pending=db.execute("SELECT count(*) FROM sources WHERE job_id=? AND state IN ('pending','processing')",(job_id,)).fetchone()[0]
            if pending:issues.append(dict(kind='processing',detail='Extraction still pending',cell_key=None))
        path=settings.data_dir/job_id/'exports'/f'revision-{job["revision"]}.xlsx'
        report=generate(job['template'],path,job['year'],job['month'],records,issues,job['revision'],job['mode'],bool(job['coverage_ack']),source_report=source_report)
        named=path.with_name(f"{job['year']}-{job['month']:02}-r{job['revision']}.xlsx")
        path.replace(named)
        with transaction() as db:
            db.execute('INSERT INTO exports(job_id,revision,path,checksum,draft,created,report) VALUES(?,?,?,?,?,?,?)',(job_id,job['revision'],str(named),sha(named.read_bytes()),int(report['draft']),time.time(),json.dumps(report)))
            db.execute("UPDATE jobs SET exported_revision=?,state=?,diagnostic=NULL WHERE id=? AND revision=? AND state NOT IN ('paused','cancelled','waiting_for_ollama')",(job['revision'],'completed_with_issues' if report['draft'] else 'completed',job_id,job['revision']))
        return query('SELECT * FROM exports WHERE job_id=? AND revision=?',(job_id,job['revision']),True)


def recalculate(export_id):
    if not settings.libreoffice:raise ValueError('Set LIBREOFFICE to the headless executable path')
    item=query('SELECT * FROM exports WHERE id=?',(export_id,),True)
    if not item:raise ValueError('Unknown export')
    with FileLock(settings.data_dir/item['job_id']/'export.lock'):
        with tempfile.TemporaryDirectory() as td:
            directory=Path(td); profile=(directory/'profile').as_uri(); out=directory/'out';out.mkdir()
            subprocess.run([settings.libreoffice,f'-env:UserInstallation={profile}','--headless','--convert-to','xlsx','--outdir',str(out),item['path']],check=True,timeout=300,capture_output=True)
            result=out/Path(item['path']).name
            from openpyxl import load_workbook
            wb=load_workbook(result,data_only=True)
            from .layout import SCHEDULE, LEGACY_EXCEL_COLS, month_cell, dates
            report=json.loads(item['report'])
            job=query('SELECT year,month FROM jobs WHERE id=?',(item['job_id'],),True)
            try:
                monthly={}
                if report.get('layout')=='monthly':
                    monthly={c.coordinate:c.value for row in wb[report['entry_sheet']].iter_rows() for c in row}
                for tc,average in report['expected_averages'].items():
                    t,c=tc.split('|')
                    if report.get('layout')=='monthly':
                        nums=[monthly.get(month_cell(d.isoformat(),t,c)) for d in dates(job['year'],job['month'])]
                        nums=[n for n in nums if isinstance(n,(int,float))]
                        actual=sum(nums)/len(nums) if nums else None
                    else:actual=wb['MASTER'][f'{LEGACY_EXCEL_COLS[c]}{SCHEDULE[t]}'].value
                    if not isinstance(actual,(int,float)) or abs(actual-average)>1e-8:raise ValueError('Recalculated workbook verification failed')
            finally:wb.close()
            # Keep the faithful original and expose recalculated artifact separately.
            target=Path(item['path']).with_suffix('.recalculated.xlsx'); atomic_write(target,result.read_bytes())
            with transaction() as db:
                report=json.loads(item['report']); report['recalculated_path']=str(target);report['recalculated_checksum']=sha(result.read_bytes())
                db.execute('UPDATE exports SET recalculated=1,report=? WHERE id=?',(json.dumps(report),export_id))
            return target
