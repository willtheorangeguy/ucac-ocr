import io
import zipfile
from pathlib import Path
import pytest
from openpyxl import load_workbook
from app.workbook import parts,sheets,parse,xml,set_cell,generate
from app.ingestion import create_job
from app.db import query
from app.exporting import export_job


def test_continuation_preserves_existing_and_rejects_other_month(data,tmp_path):
    fixture=Path(__file__).parent/'fixtures/Head Counts.xlsx'
    p=parts(fixture);_,_,ss=sheets(p);sheet=ss['September 9'][1]
    root=parse(p[sheet]);set_cell(root,'C10',5);p[sheet]=xml(root)
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w',zipfile.ZIP_DEFLATED) as z:
        for name,body in p.items():z.writestr(name,body)
    template=tmp_path/'populated.xlsx';template.write_bytes(b.getvalue())
    with pytest.raises(ValueError):create_job(2026,10,b.getvalue(),'continuation',[('bad.jpg',b'bad')])
    job=create_job(2026,9,b.getvalue(),'continuation',[('bad.jpg',b'bad')])
    assert query('SELECT value FROM cells WHERE job_id=?',(job,),True)['value']==5
    exported=export_job(job)
    wb=load_workbook(exported['path'],read_only=True)
    assert wb['DATA ENTRY - SEP 2026']['D564'].value==5
    assert 'MASTER' not in wb.sheetnames
    assert 'October 1' not in wb.sheetnames
    wb.close()
