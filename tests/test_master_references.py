from pathlib import Path
import pytest
from openpyxl import load_workbook
from app.workbook import master_reference,validate_template,generate,parts,sheets,parse,S


@pytest.mark.parametrize('formula',[
    '=IFERROR(AVERAGE(\'Boundary B\'!C10),"")',
    '=IFERROR(AVERAGE(\'September 9:Boundary B\'!C10),"")',
    '=IFERROR(AVERAGE(\'Boundary B\'!$C$10),"")',
    '= iferror ( average ( \'Boundary B\'!C10 ), "" ) ',
])
def test_valid_master_references(formula):
    assert master_reference(formula,'C10')


@pytest.mark.parametrize('formula',[
    '=IFERROR(SUM(\'Boundary B\'!C10),"")',
    '=IFERROR(AVERAGE(\'Boundary B\'!C12),"")',
    '=IFERROR(AVERAGE(\'September 9:MASTER\'!C10),"")',
    '=IFERROR(AVERAGE(\'Boundary B\'!C10),0)',
    '=IFERROR(AVERAGE(\'Boundary B\'!C10:C12),"")',
    None,
])
def test_invalid_master_references(formula):
    assert master_reference(formula,'C10') is None


def test_actual_uploaded_boundary_only_template(tmp_path):
    template=Path(__file__).parent.parent/'validation/regression/boundary-only-template.xlsx'
    assert validate_template(template)==[]
    out=tmp_path/'month.xlsx'
    record=dict(date='2026-10-09',time='05:15',column_name='public',value=0,status='accepted',provenance='synthetic')
    generate(template,out,2026,10,[record],[],1)
    wb=load_workbook(out,read_only=True)
    assert wb['DATA ENTRY - OCT 2026']['D564'].value==0
    assert 'MASTER' not in wb.sheetnames
    wb.close()



def test_web_upload_accepts_actual_template(data,photo):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.db import query
    template=Path(__file__).parent.parent/'validation/regression/boundary-only-template.xlsx'
    with TestClient(app) as client:
        response=client.post('/jobs',data={'month':'2026-10','mode':'new-month'},
            files=[('template',('Head Counts.xlsx',template.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')),
                   ('photos',('photo.png',photo,'image/png'))],follow_redirects=False)
        assert response.status_code==303, response.text
        job=response.headers['location'].split('/')[-1]
        assert query('SELECT id FROM jobs WHERE id=?',(job,),True)
        assert len(query('SELECT * FROM sources WHERE job_id=?',(job,)))==1
