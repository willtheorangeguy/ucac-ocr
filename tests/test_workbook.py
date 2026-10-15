from pathlib import Path
import zipfile
import pytest
from openpyxl import load_workbook
from app.workbook import generate, validate_template, parts, sheets, parse, xml, S
from app.monthly_workbook import DEFAULT_TEMPLATE, cells, value, strings
from app.layout import SCHEDULE, month_cell

TEMPLATE=Path(__file__).parent/'fixtures/Head Counts.xlsx'


def record(day,t,col,val):
    return dict(date=day,time=t,column_name=col,value=val,status='accepted',provenance='synthetic')


def read_entries(path,name='DATA ENTRY - OCT 2026'):
    p=parts(path);_,_,ss=sheets(p)
    return cells(parse(p[ss[name][1]])),strings(p)


@pytest.mark.parametrize('template',[TEMPLATE,DEFAULT_TEMPLATE])
def test_both_formats_export_only_monthly(template,tmp_path):
    records=[record('2026-10-09',t,c,v) for t in SCHEDULE for c,v in [('public',0),('group',32),('dry',None)]]
    out=tmp_path/'October.xlsx'
    report=generate(template,out,2026,10,records,[],1)
    assert report['draft'] and report['layout']=='monthly'
    p=parts(out);_,_,ss=sheets(p)
    assert list(ss)==['BLANK','DATA ENTRY - OCT 2026','IMPORT REPORT']
    cs,shared=read_entries(out)
    for t in SCHEDULE:
        assert value(cs[month_cell('2026-10-09',t,'public')],shared)==0
        assert value(cs[month_cell('2026-10-09',t,'group')],shared)==32
        assert value(cs.get(month_cell('2026-10-09',t,'dry')),shared) is None
        assert value(cs.get(month_cell('2026-10-09',t,'external')),shared) is None
    original=parts(DEFAULT_TEMPLATE)
    for name,data in original.items():
        if name=='xl/styles.xml' or '/printerSettings/' in name:assert p[name]==data
    blank=parse(original[sheets(original)[2]['BLANK'][1]])
    entry=parse(p[ss['DATA ENTRY - OCT 2026'][1]])
    for tag in ['cols','mergeCells','sheetProtection','dataValidations','pageMargins','pageSetup']:
        from lxml import etree as E
        assert E.tostring(blank.find('{'+S+'}'+tag))==E.tostring(entry.find('{'+S+'}'+tag))
    wb=load_workbook(out,read_only=True)
    assert wb.active.title=='DATA ENTRY - OCT 2026'
    assert wb['DATA ENTRY - OCT 2026']['G562'].value.date().isoformat()=='2026-10-09'
    assert wb['DATA ENTRY - OCT 2026']['C595'].value.strftime('%H:%M')=='13:00'
    wb.close()


def test_leap_year_blanks_zeros_averages(tmp_path):
    out=tmp_path/'February.xlsx'
    report=generate(TEMPLATE,out,2024,2,[record('2024-02-01','05:15','public',0),record('2024-02-02','05:15','public',32),record('2024-02-03','05:15','public',None)],[],2,coverage_ack=True)
    assert not report['draft']
    assert report['expected_averages']['05:15|public']==16
    wb=load_workbook(out,read_only=True);sheet=wb['DATA ENTRY - FEB 2024']
    assert sheet['G1962'].value.day==29
    assert sheet['G2032'].value is None
    assert sheet['D144'].value is None
    wb.close()


def test_modern_continuation_preserves_external_and_notes(tmp_path):
    first=tmp_path/'first.xlsx';second=tmp_path/'second.xlsx'
    generate(DEFAULT_TEMPLATE,first,2026,10,[record('2026-10-31','21:45','external',17)],[],1)
    p=parts(first);_,_,ss=sheets(p);path=ss['DATA ENTRY - OCT 2026'][1]
    root=parse(p[path]);c=cells(root)['H4'];c.set('t','inlineStr')
    from lxml import etree as E
    E.SubElement(E.SubElement(c,'{'+S+'}is'),'{'+S+'}t').text='Keep this note'
    p[path]=xml(root)
    with zipfile.ZipFile(first,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in p.items():z.writestr(n,b)
    daily=validate_template(first,2026,10)
    assert daily[-1]['counts']['21:45|external']==17
    generate(first,second,2026,10,[record('2026-10-09','05:15','group',32)],[],2,mode='continuation')
    cs,shared=read_entries(second)
    assert value(cs['E2170'],shared)==17
    assert value(cs['F564'],shared)==32
    assert value(cs['H4'],shared)=='Keep this note'
    reset=tmp_path/'reset.xlsx'
    generate(first,reset,2026,10,[],[],3,mode='new-month')
    cs,shared=read_entries(reset)
    assert value(cs['E2170'],shared) is None
    assert value(cs['H4'],shared) is None


def test_monthly_continuation_imports_each_count_column(data,tmp_path):
    from app.ingestion import create_job
    from app.db import query
    template=tmp_path/'modern.xlsx'
    records=[record('2026-10-09','05:15',c,v) for c,v in [('public',0),('external',11),('group',32),('dry',2)]]
    generate(DEFAULT_TEMPLATE,template,2026,10,records,[],1)
    job=create_job(2026,10,template.read_bytes(),'continuation',[('bad.jpg',b'bad')])
    stored={c['column_name']:c['value'] for c in query('SELECT * FROM cells WHERE job_id=?',(job,))}
    assert stored==dict(public=0,external=11,group=32,dry=2)


def test_legacy_export_upgrade_is_once_only_and_preserves_artifacts(data):
    from conftest import make_job
    from app.db import query,transaction,migrate
    from app.exporting import export_job
    import time
    job=make_job(data);old=data/job/'exports/old.xlsx';old.parent.mkdir();old.write_bytes(TEMPLATE.read_bytes())
    original=old.read_bytes()
    with transaction() as db:
        db.execute('INSERT INTO exports(job_id,revision,path,checksum,draft,created,report) VALUES(?,?,?,?,?,?,?)',(job,1,str(old),'old',1,time.time(),'{}'))
        db.execute("UPDATE jobs SET revision=1,exported_revision=1,state='completed_with_issues' WHERE id=?",(job,))
    migrate();migrate()
    assert query('SELECT revision FROM jobs WHERE id=?',(job,),True)['revision']==2
    result=export_job(job)
    assert result['revision']==2
    assert __import__('json').loads(result['report'])['layout']=='monthly'
    migrate()
    assert query('SELECT revision FROM jobs WHERE id=?',(job,),True)['revision']==2
    assert old.read_bytes()==original
    assert query('SELECT count(*) AS n FROM attempts',one=True)['n']==0


def test_actual_fall_workbook_preserves_other_months_and_summaries(tmp_path):
    uploaded=Path(__file__).parent.parent/'Head Counts - FALL 2026.xlsx'
    if not uploaded.exists():pytest.skip('Local operator workbook is not part of the repository')
    out=tmp_path/'fall.xlsx'
    records=[record('2026-10-01','05:15',c,v) for c,v in [('public',0),('external',11),('group',32),('dry',2)]]
    generate(uploaded,out,2026,10,records,[],1)
    p=parts(uploaded);q=parts(out);_,_,ss=sheets(p)
    assert q[ss['DATA ENTRY - SEP 2026'][1]]==p[ss['DATA ENTRY - SEP 2026'][1]]
    assert q[ss['BLANK'][1]]==p[ss['BLANK'][1]]
    assert q['xl/styles.xml']==p['xl/styles.xml']
    cs,shared=read_entries(out)
    assert [value(cs[c+'4'],shared) for c in 'DEFG']==[0,11,32,2]
    summary=parse(q[ss['OCT 2026'][1]])
    assert 'DATA ENTRY - OCT 2026' in cells(summary)['C3'].find('{'+S+'}f').text
    assert len(validate_template(out,2026,10))==31
