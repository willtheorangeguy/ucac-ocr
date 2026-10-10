import json
import uuid
from app.db import query,transaction
from app.reconciliation import reconcile,decide
from app.services import configure_source,control
from app.exporting import export_job
from conftest import make_job,result
from test_worker import add_tasks


def success_all():
    with transaction() as db:
        for t in db.execute('SELECT * FROM tasks').fetchall():
            db.execute("UPDATE tasks SET state='succeeded',result=? WHERE id=?",(json.dumps(result(json.loads(t['expected']))),t['id']))


def test_conflict_review_protected_new_evidence(data,photo):
    job=make_job(data);add_tasks(data,job,photo,2);success_all()
    sources=query('SELECT * FROM sources ORDER BY rowid')
    reconcile(sources[0]['id']);reconcile(sources[1]['id'])
    assert query('SELECT count(*) AS n FROM issues WHERE resolved=0',one=True)['n']==0
    with transaction() as db:
        for task in db.execute('SELECT * FROM tasks WHERE source_id=?',(sources[1]['id'],)).fetchall():
            r=json.loads(task['result']);r['rows'][0]['group'].update(raw_text='30 + 2',kind='sum',terms=[30,2])
            db.execute('UPDATE tasks SET result=? WHERE id=?',(json.dumps(r),task['id']))
    reconcile(sources[1]['id'])
    cell=query('SELECT * FROM cells WHERE column_name="group"',one=True)
    assert cell['status']=='needs_review' and cell['value'] is None
    issue=query('SELECT * FROM issues WHERE resolved=0',one=True)
    decide(job,issue['id'],'value',32,'Ink checked','Reviewer')
    reconcile(sources[1]['id'])
    cell=query('SELECT * FROM cells WHERE column_name="group"',one=True)
    assert cell['reviewed'] and cell['value']==32
    assert query('SELECT * FROM issues WHERE resolved=0',one=True) is None
    with transaction() as db:db.execute('UPDATE sources SET version=version+1 WHERE id=?',(sources[1]['id'],));db.execute('UPDATE tasks SET version=version+1 WHERE source_id=?',(sources[1]['id'],))
    reconcile(sources[1]['id'])
    assert query('SELECT * FROM issues WHERE resolved=0',one=True)['kind']=='reviewed_new_evidence'


def test_date_repair_reuses_successful_tasks(data,photo):
    job=make_job(data);add_tasks(data,job,photo);success_all()
    source=query('SELECT * FROM sources',one=True)
    with transaction() as db:db.execute('UPDATE sources SET assigned_date=NULL WHERE id=?',(source['id'],))
    reconcile(source['id'])
    assert query('SELECT * FROM issues WHERE resolved=0',one=True)['kind']=='date'
    configure_source(job,source['id'],'2026-10-09',reviewer='R',confirm=True)
    assert len(query('SELECT * FROM cells'))==3
    assert len(query("SELECT * FROM tasks WHERE state='succeeded'"))==2


def test_immutable_export_revision_and_review_regeneration(data,photo):
    job=make_job(data);add_tasks(data,job,photo);success_all()
    source=query('SELECT * FROM sources',one=True);reconcile(source['id'])
    control(job,'coverage')
    first=export_job(job)
    assert not first['draft']
    with transaction() as db:
        from app.db import issue
        c=db.execute('SELECT * FROM cells WHERE column_name="public"').fetchone()
        issue(db,job,'manual:'+c['key'],'cell',{},source['id'],c['key'])
        issue_id=db.execute('SELECT id FROM issues WHERE key=?',('manual:'+c['key'],)).fetchone()[0]
    decide(job,issue_id,'blank',None,'Confirmed blank','R')
    second=export_job(job)
    assert second['revision']>first['revision']
    from openpyxl import load_workbook
    old=load_workbook(first['path'],read_only=True);new=load_workbook(second['path'],read_only=True)
    assert old['October 9']['C10'].value==0
    assert new['October 9']['C10'].value is None
    old.close();new.close()
